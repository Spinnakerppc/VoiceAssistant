"""
nas_search.py — NAS document search for CannaKit Voice Assistant
"""
import os
import logging
import re
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)

NAS_MOUNT_PATH    = "/mnt/nas"   # ← update to your actual mount path
MAX_SCAN_FILES    = 2000
MAX_RESULTS       = 5
MAX_CONTENT_BYTES = 1_000_000

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".csv"}


def _extract_txt(path):
    try:
        return path.read_text(errors="ignore")[:MAX_CONTENT_BYTES]
    except Exception:
        return ""

def _extract_csv(path):
    try:
        return path.read_text(errors="ignore")[:MAX_CONTENT_BYTES]
    except Exception:
        return ""

def _extract_pdf(path):
    try:
        import pdfplumber
        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    parts.append(t)
                if sum(len(p) for p in parts) >= MAX_CONTENT_BYTES:
                    break
        return "\n".join(parts)
    except ImportError:
        return ""
    except Exception:
        return ""

def _extract_docx(path):
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)[:MAX_CONTENT_BYTES]
    except ImportError:
        return ""
    except Exception:
        return ""

def _extract_xlsx(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join(str(c) for c in row if c is not None)
                if row_text.strip():
                    parts.append(row_text)
        return "\n".join(parts)[:MAX_CONTENT_BYTES]
    except ImportError:
        return ""
    except Exception:
        return ""

CONTENT_EXTRACTORS = {
    ".txt":  _extract_txt,
    ".csv":  _extract_csv,
    ".pdf":  _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def _nas_available():
    return os.path.isdir(NAS_MOUNT_PATH) and os.access(NAS_MOUNT_PATH, os.R_OK)


def _make_snippet(text, keyword, context=60):
    idx = text.lower().find(keyword.lower())
    if idx == -1:
        return ""
    start = max(0, idx - context)
    end   = min(len(text), idx + len(keyword) + context)
    snippet = text[start:end].replace("\n", " ").strip()
    if start > 0:
        snippet = "…" + snippet
    if end < len(text):
        snippet = snippet + "…"
    return snippet


def search_nas(query, content_search=True):
    if not _nas_available():
        log.warning(f"[NAS] Mount not available at {NAS_MOUNT_PATH}")
        return []
    keywords = [w.lower() for w in re.split(r"\s+", query.strip()) if w]
    if not keywords:
        return []
    results = []
    scanned = 0
    for root, dirs, files in os.walk(NAS_MOUNT_PATH):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for filename in files:
            if scanned >= MAX_SCAN_FILES:
                return results
            ext = Path(filename).suffix.lower()
            if ext not in SUPPORTED_EXTENSIONS:
                continue
            scanned += 1
            filepath = Path(root) / filename
            name_lower = filename.lower()
            if all(kw in name_lower for kw in keywords):
                results.append({"name": filename, "path": str(filepath), "match": "filename", "snippet": ""})
                if len(results) >= MAX_RESULTS:
                    return results
                continue
            if content_search:
                extractor = CONTENT_EXTRACTORS.get(ext)
                if not extractor:
                    continue
                content = extractor(filepath)
                if all(kw in content.lower() for kw in keywords):
                    results.append({"name": filename, "path": str(filepath), "match": "content", "snippet": _make_snippet(content, keywords[0])})
                    if len(results) >= MAX_RESULTS:
                        return results
    return results


def _extract_search_query(text):
    import re as _re
    # Strip common filler phrases first
    text = text.lower().strip().rstrip("?.")
    fillers = [
        r"any file[s]? on (?:the )?(?:nas|network attached storage|network|storage)[\w\s]*that contains? (?:the )?word\s+",
        r"any file[s]? on (?:the )?(?:nas|network attached storage|network|storage)[\w\s]*",
        r"that contains? (?:the )?word\s+",
        r"on (?:the )?(?:nas|network|storage)",
        r"(?:find me|find|search for|look for|locate|search)\s+",
        r"(?:the |a |my )",
    ]
    for f in fillers:
        text = _re.sub(f, " ", text).strip()
    text = " ".join(text.split())  # collapse whitespace
    if text:
        return text
    patterns = [
        r"(?:find|search for|look for|locate|find me|search)\s+(?:the\s+|a\s+|my\s+)?(.+)",
        r"(?:do you have|is there)\s+(?:a\s+|the\s+)?(?:file called|document called|file named)?\s*(.+)",
        r"(?:open|show me)\s+(?:the\s+|a\s+)?(?:file\s+)?(.+)",
    ]
    text_lower = text.lower().strip()
    for pattern in patterns:
        m = re.search(pattern, text_lower)
        if m:
            return m.group(1).strip().rstrip("?.")
    return None


def handle_nas_search(text):
    query = _extract_search_query(text)
    if not query:
        return "What would you like me to search for on the NAS?"
    log.info(f"[NAS] Search query: {query!r}")
    if not _nas_available():
        return f"I can't reach the NAS right now. Make sure it's mounted at {NAS_MOUNT_PATH}."
    results = search_nas(query)
    if not results:
        return f"I didn't find any files matching '{query}' on the NAS."
    if len(results) == 1:
        r = results[0]
        if r["match"] == "filename":
            return f"I found one file: {r['name']}, located at {r['path']}."
        snippet_part = f" It contains: {r['snippet']}" if r["snippet"] else ""
        return f"I found one file: {r['name']}.{snippet_part} It's at {r['path']}."
    names = ", ".join(r["name"] for r in results[:3])
    extra = f" and {len(results) - 3} more" if len(results) > 3 else ""
    return f"I found {len(results)} files matching '{query}': {names}{extra}. Check the log for full paths."
